from pathlib import Path
import json,html
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,PageBreak
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PInches,Pt
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parents[1];D=R/'docs';audit=json.loads((R/'data/audit.json').read_text());pred=json.loads((R/'data/prediction_report.json').read_text());tests=json.loads((R/'data/test_results.json').read_text());ui=json.loads((R/'data/ui_results.json').read_text())
sections=[
('Executive summary',[
'Boundary is an integrated capstone for franchise analysts and scouts. It combines uploaded match evidence, descriptive scouting, exact XI planning, research forecasts and a DeepSeek tool-calling analyst in a custom research workspace.',
f'The processed corpus contains {audit["matches"]:,} eligible men’s matches and {audit["match_player_rows"]:,} named player-match records. The supplied squads contain ten franchises. Historical auction records end in 2022.',
'Completed validation: 37 offline behavioral checks, 12 Streamlit workflow checks, and rendered desktop/mobile inspection. These results do not establish live language-model factuality or expert tactical usefulness.',
'Primary contribution: transparent decision support with evidence traceability and explicit limits. No claim of guaranteed predictive advantage, complete official eligibility certification or current market valuation is made.'
]),
('Problem, users and product strategy',[
'The initial app mixed a historical benchmark with a different live prediction formula, treated fixed dashboard cards as agents and omitted bench selections from lineup calculations. The integrated release addresses those concrete defects.',
'Primary user journeys: discover role-fit candidates; inspect player evidence; construct an XI after exclusions; explore a budget scenario; request a sourced brief; preserve the reasoning in a journal.',
'Strengths: uploaded multi-league history, player IDs and numerical decision tools. Weaknesses: old auction records, unverified availability and heuristic scouting weights. Opportunities: auditable scouting and expert-reviewed decisions. Threats: stale data, identity aliases, unsupported AI text and public API costs.',
'Each requested planning dimension appears in the product: SWOT, intelligence, reliability, repeatability, predictive research, qualitative review, quantitative evaluation and custom UI. Human review and live provider quality remain evidence gaps rather than fabricated successes.'
]),
('Data engineering and provenance',[
'Inputs: IPL, BBL, SMAT, Hundred and T20 international JSON ZIPs; original supplied squad CSV; supplied IPL auction CSV. Women’s matches, duplicate match IDs and the voided IPL match 1473495 are excluded from the men’s scouting corpus.',
'Each named player has one row per match. Registry IDs link records; supplied squad names normalize display aliases. Missing IDs fall back to a league/name label, so unresolved identity remains visible. The archive manifest records full SHA-256 fingerprints and date coverage.',
'Runs faced, legal deliveries, wides/no-balls, bowler-attributable runs and wicket types are computed separately. Run outs count as batter dismissals but not bowler wickets. Super overs are omitted from player regulation statistics.',
'All as-of scouting queries use actual match dates strictly before the cutoff. The fixed squad snapshot is not a reconstruction of historical membership. Full-format phase definitions are used rather than claiming rain-adjusted official rules. Hundred economy is displayed per six-ball equivalent.'
]),
('Scouting, uncertainty and exact planning',[
'For a selected evidence pool, scouting uses a 120-ball prior at the pool rate. Batting score is the shrunk difference from pool runs per ball multiplied by 100. Bowling score is the pool-minus-shrunk conceded rate multiplied by six. Sample weight is n/(n+120), not a probability of future success.',
'Combined competitions are not strength-adjusted. This release does not label descriptive scores as RAA or forecasts. Candidate comparisons expose raw rates and sample sizes. Observed-rate uncertainty uses 400 reproducible match-bootstrap samples; independent-match assumptions and scope are disclosed.',
'The XI solver maximizes a role-weighted utility: batting score/10 for batting roles plus bowling score for bowling roles. Missing evidence receives -1. Constraints: exactly 11 players, at most four overseas, at least one keeper, configurable batting/bowling coverage, user locks and exclusions.',
'Optional budget constraints use editable user scenario prices. Placeholder costs are explicitly labeled and historical auction amounts are not silently converted into current values. The budget applies to the XI only. Impact Player certification and full auction-squad optimization are not claimed.'
]),
('Prediction design and results',[
'The pre-toss model uses prior team Elo and exponentially updated result form. Post-toss adds chasing direction and prior venue chase rate. Models have no intercept to preserve team-order symmetry. Regularized logistic fitting uses IPL matches before 2023; validation is 2023; the untouched test begins in 2024.',
'All matches on a date are forecast before any outcome from that date updates state. One feature function is used for both evaluation and scenario serving. Fixed constants and regularization were not tuned on the test set. Current forecasts use the latest uploaded IPL ratings; live earlier as-of forecasts are blocked.',
'On 215 test matches: pre-toss accuracy 51.16%, Brier 0.250995, log loss 0.695140. Post-toss accuracy 54.88%, Brier 0.248129, log loss 0.689425. Post-toss Wilson 95% accuracy interval: 48.20–61.39%. Constant-50% forecast Brier: 0.250000; Elo baseline Brier: 0.258632.',
'The accuracy interval includes chance. The small Brier improvement does not establish a durable predictive advantage. Wilson intervals assume independent outcomes and omit season dependence. Lineups, weather, pitch reports and injuries are not model features. Calibration and per-season metrics are included in the app and JSON report.'
]),
('DeepSeek analyst and reliability controls',[
'The server calls DeepSeek Chat through its HTTPS API. The model selects among four approved tool types: scouting, player dossier, selected-franchise XI optimization and research match forecasts. It receives scoped evidence and produces a brief with source markers and an inspectable trace.',
'Tool arguments are validated; no arbitrary code execution, filesystem path tool or URL-fetch tool is exposed. A final answer is rejected if the model fails to retrieve evidence within the bounded workflow. Model prose is not verified claim by claim, so source markers alone are not a factuality guarantee.',
'Each brief allows three provider requests, four tools per enabled turn and 1,100 output tokens per request. A persistent local SQLite ledger caps provider requests at thirty per UTC day. Cache keys include data provenance, squad fingerprints, model weights and question settings.',
'The user-requested key is configured only in server secrets in the private delivery bundle, ignored by Git and excluded from prompt text, journals and screenshots. Provider response bodies are suppressed in errors. Outbound policy blocks the provider host in the build environment; live authentication and answer quality were not verified.'
]),
('Evaluation, UI and qualitative review',[
'37 offline checks cover parser edge cases, identity uniqueness, date cutoffs, sample filters, phase scope, all ten franchise solutions, bench scoring, locked/excluded players, infeasible budgets, forecast symmetry and temporal invariance, typed tool restrictions, caching, local call allowance and journal persistence.',
'12 headless UI checks cover eight workspaces, the saved demonstration, journal save/read, budget mode and infeasible scenarios. Real Chromium rendering was inspected at desktop and mobile sizes. The interface uses editorial typography, restrained green accents, research tables and visible evidence rather than decorative agent activity.',
'Qualitative evaluation is implemented as a reviewer form covering evidence support, cricket relevance, tradeoff clarity, usability and appropriate uncertainty. Ratings are saved only after a reviewer submits a named task. No expert ratings have been submitted or fabricated.',
'Live AI evaluation should include unknown players, unsupported prices, ambiguous roles, weak samples, impossible constraints, adversarial source instructions and provider failure. Current automated mocks validate orchestration and error boundaries, not the model’s real behavior under every attack.'
]),
('Operation, reproducibility and conclusions',[
'Run start.sh or start.bat, or install requirements and launch streamlit run app.py. Processed data and model artifacts are included. Rebuild with match_data.py over the original archives, then prediction.py, tests/test_core.py and tests/test_ui.py.',
'No paid data feed, GPU, domain or remote database is required. A small DeepSeek allowance can support a controlled capstone demonstration. Actual pricing and input-token costs must be checked with the provider. Local limits do not replace provider billing controls or authentication for a public site.',
'Decisions, reviews and usage live in .runtime, excluded from delivery. Free hosting may be ephemeral; export important journal records. Multitenant access control, a remote durable database and cross-replica spend accounting are not implemented. Source-data terms must be checked before public redistribution.',
'The integrated release makes the analytical and engineering work concrete and reviewable. It supports useful, reproducible research decisions while preserving the distinction between evidence, descriptive estimates and recommendations. Remaining evaluation gaps are stated explicitly.'
])]
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodyBoundary',fontSize=10,leading=15,spaceAfter=12,textColor=colors.HexColor('#273c3e')));styles['Title'].fontName='Times-Roman';styles['Title'].fontSize=30;styles['Title'].leading=33;styles['Heading1'].textColor=colors.HexColor('#176b55')
story=[Paragraph('BOUNDARY',styles['Title']),Paragraph('Franchise research studio<br/>Integrated capstone report · 18 October 2026',styles['Heading2']),Spacer(1,20)]
for i,(title,paragraphs) in enumerate(sections):
 if i:story.append(PageBreak())
 story.append(Paragraph(str(i+1)+'. '+title,styles['Heading1']))
 for para in paragraphs:story.append(Paragraph(html.escape(para),styles['BodyBoundary']))
 if i==0:story.append(Image(str(D/'previews/decision-desk.png'),width=460,height=351))
 if i==2:
  data=[['Competition','Included men’s matches','Latest']]+[[x['competition'],str(x['included']),x['latest']] for x in audit['sources']];t=Table(data,colWidths=[140,170,120]);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#edf0e8')),('TEXTCOLOR',(0,0),(-1,0),colors.HexColor('#176b55')),('BOTTOMPADDING',(0,0),(-1,-1),9),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#d5ddd4'))]));story.append(t)
def footer(c,d):c.setFont('Helvetica',8);c.setFillColor(colors.HexColor('#627369'));c.drawString(42,26,'BOUNDARY / CAPSTONE RESEARCH EDITION');c.drawRightString(550,26,str(d.page))
SimpleDocTemplate(str(D/'Capstone_Report.pdf'),topMargin=38,bottomMargin=42,leftMargin=42,rightMargin=42,title='Boundary integrated capstone report').build(story,onFirstPage=footer,onLaterPages=footer)
doc=Document();doc.add_heading('Boundary — integrated capstone report',0);doc.add_paragraph('18 October 2026 · Franchise analysts and scouts')
for i,(title,paras) in enumerate(sections):
 doc.add_heading(str(i+1)+'. '+title,1)
 for para in paras:doc.add_paragraph(para)
 if i==0:doc.add_picture(str(D/'previews/decision-desk.png'),width=Inches(6))
doc.save(D/'Capstone_Report.docx')
slides=[('Boundary','Franchise research studio\nOne integrated capstone release\n18 October 2026'),('The decision problem','Find a role-fit player\nCompare evidence and uncertainty\nRebuild an XI after an exclusion\nPreserve a defensible recommendation'),('The evidence base','6,392 eligible men’s matches\n141,377 named player-match records\nIPL · BBL · SMAT · Hundred · T20I\nSource hashes and strict match-date cutoffs'),('Scouting and observed uncertainty','Pool-relative scores with 120-ball shrinkage\nPhase-specific batting and bowling comparisons\nReproducible match-bootstrap intervals\nNo claim of competition-strength translation'),('Exact XI and budget planning','11 players · overseas and role constraints\nLocks, exclusions and full-bench scoring\nEditable scenario prices and budget\nExplicit planning assumptions, not official certification'),('DeepSeek analyst','Question → approved tool calls → cited brief\nNumerical tools own calculations\nPrivate server-side key configuration\nBounded calls, trace, cache and journal'),('Prediction research','One shared evaluation and serving pipeline\nTrain before 2023 · validate 2023 · test 2024–2026\nPost-toss 54.88% accuracy on 215 matches\n95% interval 48.20–61.39%; Brier 0.248129'),('Validation and honest limits','37 offline behavioral checks passed\n12 UI workflow checks passed\nDesktop and mobile rendering inspected\nLive AI quality and expert review remain unvalidated'),('SWOT in action','Strength: uploaded history and exact local tools\nWeakness: old prices and unverified availability\nOpportunity: traceable scouting decisions\nThreat: stale data, unsupported claims and API cost'),('Demonstration and conclusion','Scout → compare → inspect dossier → solve XI\nInspect forecast calibration and source evidence\nSave a decision and collect a real review\nLow infrastructure cost; transparent decision support')]
prs=Presentation();prs.slide_width=PInches(13.333);prs.slide_height=PInches(7.5)
for index,(title,body) in enumerate(slides):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('F6F5F0')
 box=slide.shapes.add_textbox(PInches(.65),PInches(.45),PInches(12),PInches(.4));tf=box.text_frame;tf.text=f'BOUNDARY / CAPSTONE RESEARCH                         {index+1:02d}';tf.paragraphs[0].font.size=Pt(11);tf.paragraphs[0].font.color.rgb=RGBColor.from_string('176B55')
 box=slide.shapes.add_textbox(PInches(.65),PInches(1.25),PInches(12),PInches(1.1));tf=box.text_frame;tf.text=title;tf.paragraphs[0].font.name='Georgia';tf.paragraphs[0].font.size=Pt(37);tf.paragraphs[0].font.color.rgb=RGBColor.from_string('182D32')
 box=slide.shapes.add_textbox(PInches(.75),PInches(2.9),PInches(11.7),PInches(4));tf=box.text_frame;tf.word_wrap=True
 for n,line in enumerate(body.split('\n')):
  q=tf.paragraphs[0] if n==0 else tf.add_paragraph();q.text=line;q.font.size=Pt(25);q.space_after=Pt(22);q.font.color.rgb=RGBColor.from_string('344D4E')
prs.save(D/'Capstone_Presentation.pptx')
print('Created PDF, editable DOCX and 10-slide presentation.')
